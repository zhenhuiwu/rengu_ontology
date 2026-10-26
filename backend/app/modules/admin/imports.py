"""用户模板解析：CSV 向后兼容，XLSX 交给 openpyxl；输入有界且不落盘。"""

import base64
import binascii
import csv
import io
import zlib
from xml.etree.ElementTree import ParseError
from zipfile import ZIP_DEFLATED, ZIP_STORED, BadZipFile, ZipFile

from defusedxml.common import DefusedXmlException
from openpyxl import load_workbook
from openpyxl.packaging.manifest import Manifest
from openpyxl.utils.exceptions import InvalidFileException
from openpyxl.xml.constants import XLSX
from openpyxl.xml.functions import fromstring
from pydantic import ValidationError

from app.core.errors import AppError
from app.modules.admin.schemas import IMPORT_MAX_ROWS, IMPORT_XLSX_MAX_BYTES, ImportRowIn

IMPORT_HEADER = {"姓名": "name", "账号": "account", "部门": "dept", "角色": "role"}
XLSX_MAX_EXPANDED = 20 * 1024 * 1024
XLSX_MAX_ENTRIES = 128
XLSX_MAX_SCAN_ROWS = 10_000
XLSX_MAX_COLUMNS = 32


def invalid(message: str) -> AppError:
    return AppError(message, code="VALIDATION_FAILED", status_code=422)


def header_keys(header: list[str], *, required: bool = False) -> list[str | None]:
    keys = [IMPORT_HEADER.get(h.strip()) for h in header]
    known = [k for k in keys if k is not None]
    if len(known) != len(set(known)):
        raise invalid("模板包含重复的姓名、账号、部门或角色表头")
    if required and set(known) != set(IMPORT_HEADER.values()):
        raise invalid("Excel 首行必须包含姓名、账号、部门、角色表头")
    return keys


def append_row(out: list[ImportRowIn], keys: list[str | None], line: list[str]) -> None:
    if not any(c.strip() for c in line):
        return
    if len(out) >= IMPORT_MAX_ROWS:
        raise invalid("一次最多导入 500 人，请拆分文件")
    out.append(ImportRowIn(**{k: line[i].strip() if i < len(line) else "" for i, k in enumerate(keys) if k}))


def parse_import_csv(text: str) -> list[ImportRowIn]:
    """保留 BOM、列重排、缺列为空、忽略未知列及空行；重复已知表头不再覆盖。"""
    reader = csv.reader(io.StringIO(text.lstrip("\ufeff")))
    try:
        header = next(reader, None)
        if not header:
            return []
        keys = header_keys(header)
        out: list[ImportRowIn] = []
        for line in reader:
            append_row(out, keys, line)
        return out
    except (csv.Error, ValidationError) as exc:
        raise invalid("CSV 内容无效，请检查模板和文本字符") from exc


def checked_archive(data: bytes) -> None:
    """只检查 ZIP 容器，不自行解析 OOXML；不解压到文件系统。"""
    with ZipFile(io.BytesIO(data)) as archive:
        members = archive.infolist()
        if len(members) > XLSX_MAX_ENTRIES or sum(m.file_size for m in members) > XLSX_MAX_EXPANDED:
            raise invalid("Excel 压缩包过大或包含过多文件")
        names = [m.filename.replace("\\", "/").lower() for m in members]
        if len(names) != len(set(names)):
            raise invalid("Excel 压缩包包含重复文件")
        if any("vbaproject" in name or "macrosheets/" in name for name in names):
            raise invalid("不支持带宏的 Excel，请另存为无宏 .xlsx")
        expanded = 0
        for member in members:
            if member.flag_bits & 1 or member.compress_type not in (ZIP_STORED, ZIP_DEFLATED):
                raise invalid("Excel 不支持加密或特殊压缩格式")
            with archive.open(member) as entry:
                while chunk := entry.read(64 * 1024):
                    expanded += len(chunk)
                    if expanded > XLSX_MAX_EXPANDED:
                        raise invalid("Excel 解压内容超过 20 MiB")
        # 复用 openpyxl 的包清单模型识别类型，不能靠上传文件扩展名判断宏。
        manifest = Manifest.from_tree(fromstring(archive.read("[Content_Types].xml")))
        content_types = [p.ContentType for p in [*manifest.Default, *manifest.Override]]
        if XLSX not in content_types or any(
            "macroenabled" in t.lower() or "vbaproject" in t.lower() for t in content_types
        ):
            raise invalid("仅支持无宏 .xlsx 文件")


def parse_import_xlsx(encoded: str) -> list[ImportRowIn]:
    try:
        data = base64.b64decode(encoded, validate=True)
    except (binascii.Error, ValueError) as exc:
        raise invalid("Excel 文件编码无效") from exc
    if len(data) > IMPORT_XLSX_MAX_BYTES:
        raise invalid("Excel 文件不能超过 2 MiB")
    try:
        checked_archive(data)
        book = load_workbook(io.BytesIO(data), read_only=True, data_only=False, keep_links=False)
        try:
            if not book.worksheets:
                return []
            sheet = book.worksheets[0]
            if sheet.max_column and sheet.max_column > XLSX_MAX_COLUMNS:
                raise invalid("Excel 模板最多支持 32 列")
            # 不信任 dimension，按实际行宽校验，不能把越界列或隐藏公式静默截断。
            sheet.reset_dimensions()
            out: list[ImportRowIn] = []
            keys: list[str | None] | None = None
            for index, cells in enumerate(sheet.iter_rows(), 1):
                if index > XLSX_MAX_SCAN_ROWS:
                    raise invalid("Excel 扫描范围超过 10000 行，请移除多余空行或格式")
                if len(cells) > XLSX_MAX_COLUMNS:
                    raise invalid("Excel 模板最多支持 32 列")
                if any(c.data_type == "f" for c in cells):
                    raise invalid(f"Excel 第 {index} 行包含公式，请改为文本后导入")
                line = ["" if c.value is None else str(c.value) for c in cells]
                if keys is None:
                    if not any(c.strip() for c in line):
                        continue
                    keys = header_keys(line, required=True)
                else:
                    append_row(out, keys, line)
            return out
        finally:
            book.close()
    except (
        BadZipFile,
        zlib.error,
        InvalidFileException,
        DefusedXmlException,
        ParseError,
        ValidationError,
        ValueError,
        KeyError,
        IndexError,
        OverflowError,
        TypeError,
        OSError,
        RuntimeError,
    ) as exc:
        raise invalid("Excel 文件无效，请使用未加密、无宏的 .xlsx 模板") from exc


def row_length_error(row: ImportRowIn) -> str | None:
    for name, limit, label in (
        ("name", 64, "姓名"),
        ("account", 128, "账号"),
        ("dept", 128, "部门"),
        ("role", 64, "角色"),
    ):
        if len(getattr(row, name)) > limit:
            return f"{label}不能超过 {limit} 个字符"
    return None
